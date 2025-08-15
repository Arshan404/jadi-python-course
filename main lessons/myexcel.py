import openpyxl

wb = openpyxl.load_workbook("E:\خلاصه کتاب.xlsx")
print(type(wb))

ws = wb.active
print(ws)
print(ws["A4"])
print(ws["A4"].value)